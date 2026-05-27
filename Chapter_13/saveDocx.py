from docx import Document

doc = Document()
doc.add_paragraph('Hello, World!')

paraObj1 = doc.add_paragraph("It's second paragraph.")
paraObj2 = doc.add_paragraph("It's paragraph too.")

paraObj1.add_run(' This text will bee add in second paragraph.')

doc.save('multipleParagrahs.docx')
