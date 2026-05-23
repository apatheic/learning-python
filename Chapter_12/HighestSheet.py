from openpyxl import load_workbook

wb = load_workbook('example3.xlsx')
sheet = wb['Sheet1']
print(sheet.max_column)

print(sheet.max_row)
