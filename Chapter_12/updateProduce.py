from openpyxl import load_workbook

wb = load_workbook('produceSales.xlsx')
sheet = wb['Sheet']

#Product types and their updated prices
PRICE_UPDATES = {
        'Lemon': 3.07,
                 'Celery': 1.19,
                        'Чеснок': 1.27
}

for rowNum in range(2, sheet.max_row):

    produceName = sheet.cell(row=rowNum, column=1).value
    if produceName in PRICE_UPDATES:
        sheet.cell(row=rowNum, column=2).value = PRICE_UPDATES[produceName]

wb.save('updatedProduceSales.xlsx')
