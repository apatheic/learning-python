from openpyxl import load_workbook

wb = load_workbook('writeFormula.xlsx')
sheet = wb.active

sheet['A3'].value
#'=SUM(A1:A2)'

wbDataOnly = load_workbook('writeFormula.xlsx', 
                           data_only=True)
sheet = wbDataOnly.active
sheet['A3'].value
#500
