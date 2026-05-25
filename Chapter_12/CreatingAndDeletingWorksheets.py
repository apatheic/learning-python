from openpyxl import Workbook

#create the first sheet
wb = Workbook()
print(wb.sheetnames)

# Create an additional sheet with the automatic name "Sheet1"
wb.create_sheet()
print(wb.sheetnames)

# Create a new sheet named "First Sheet" and place it at the beginning (index=0)
wb.create_sheet(index=0, title='First Sheet')
print(wb.sheetnames)

# Create a "Middle Sheet" in position 2 (between "Sheet" and "Sheet1")
wb.create_sheet(index=2, title='Middle Sheet')
print(wb.sheetnames)

#Deleting "Sheet1"
del wb['Sheet1']

#Deleting "Middle Sheet"
del wb['Middle Sheet']
