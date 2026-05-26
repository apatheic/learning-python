from openpyxl import Workbook

wb = Workbook()
sheet = wb.active

#here we merged cells A1-D3
sheet.merge_cells('A1:D3')
sheet['A1'] = '12 cells were merged.'

#and here we merged cells C5-E5
sheet.merge_cells('C5:E5')
sheet['C5'] = '3 cells were merged.'

wb.save('merged.xlsx')
