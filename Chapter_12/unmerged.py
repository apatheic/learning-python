from openpyxl import load_workbook

wb = load_workbook('merged.xlsx')
sheet = wb.active

#here we unmerged cells A1-D3
sheet.unmerge_cells('A1:D3')

#and here we unmerged cells C5-E5
sheet.unmerge_cells('C5:E5')
