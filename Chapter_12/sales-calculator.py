from openpyxl import load_workbook, Workbook

def _get_rows_count(wb: Workbook, sheet_name: str) -> int:
    sheet = wb[sheet_name]
    return len(list(sheet.rows))

wb = load_workbook('./sales.xlsx', data_only=True)
sheet = wb['sales_04']
sheet['E1'].value = 'Sales'

for r in range(2, _get_rows_count(wb, 'sales_04') + 1):
    dt = sheet.cell(r,1).value
    if dt is None:
        break
    fruit = sheet.cell(r,2).value  # ✅ ИСПРАВИЛ: был минус
    quantity = float(sheet.cell(r,3).value)
    unit_price = float(sheet.cell(r,4).value)
    if quantity is not None and unit_price is not None:
        sales = quantity * unit_price
        sheet.cell(r,5).value = sales
        print(f'{dt} - {fruit} - {quantity} - {unit_price} - {sales}')

wb.save('./sales.xlsx')
